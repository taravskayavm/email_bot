from __future__ import annotations

from pathlib import Path
import csv
import zipfile
import io

import fitz  # PyMuPDF
from docx import Document
from openpyxl import Workbook


def make_pdf(path: Path, blocks):
    doc = fitz.open()
    page = doc.new_page()
    x, y = 50, 72
    for text, opts in blocks:
        opts = opts or {}
        if opts.get("newline"):
            page.insert_text((x, y), " ", fontsize=12)
            x += 7
            continue
        sup = bool(opts.get("superscript"))
        size = 8 if sup else 12
        dy = -4 if sup else 0
        page.insert_text((x, y + dy), text, fontsize=size)
        x += size * 0.6 * len(text) + 6
    out = io.BytesIO()
    doc.save(out)
    doc.close()
    path.write_bytes(out.getvalue())
    return path


def make_docx(tmp_path: Path, lines: list[str]) -> Path:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    path = tmp_path / "temp.docx"
    doc.save(path)
    return path


def make_xlsx(tmp_path: Path, rows: list[list[str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    path = tmp_path / "temp.xlsx"
    wb.save(path)
    return path


def make_csv(tmp_path: Path, lines: list[str]) -> Path:
    path = tmp_path / "temp.csv"
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for line in lines:
            writer.writerow([line])
    return path


def make_zip(tmp_path: Path, paths: list[Path]) -> Path:
    path = tmp_path / "temp.zip"
    with zipfile.ZipFile(path, "w") as z:
        for p in paths:
            z.write(p, arcname=p.name)
    return path
